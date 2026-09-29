#!/usr/bin/env python3
"""Build lab_sample_size.xlsx -- the Excel version of the five calculations.

Every answer cell is a live formula. Change an input in blue and the whole
sheet recomputes. Nothing is hard-coded and there are no macros.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "1F4E79"
BLUE = "2E75B6"
AMBER = "B85C22"
GREY = "595959"

H1 = Font(name="Calibri", size=14, bold=True, color=NAVY)
H2 = Font(name="Calibri", size=11, bold=True, color=NAVY)
LBL = Font(name="Calibri", size=11)
INP = Font(name="Consolas", size=11, bold=True, color=BLUE)
OUT = Font(name="Consolas", size=11, color="1F2937")
ANS = Font(name="Consolas", size=12, bold=True, color=AMBER)
NOTE = Font(name="Calibri", size=10, italic=True, color=GREY)
MONO = Font(name="Consolas", size=10, color=GREY)

FILL_IN = PatternFill("solid", fgColor="EAF2FB")
FILL_ANS = PatternFill("solid", fgColor="FFF8E1")
FILL_HDR = PatternFill("solid", fgColor="E7EEF5")
THIN = Side(style="thin", color="B8C6D4")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

wb = Workbook()


def sheet(title):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 52
    return ws


def put(ws, row, label, value=None, note=None, font=LBL, vfont=OUT,
        fill=None, fmt=None):
    ws.cell(row, 2, label).font = font
    if value is not None:
        c = ws.cell(row, 3, value)
        c.font = vfont
        c.alignment = Alignment(horizontal="right")
        if fill:
            c.fill = fill
            c.border = BOX
        if fmt:
            c.number_format = fmt
    if note:
        ws.cell(row, 4, note).font = NOTE
    return row + 1


def title_block(ws, n, heading, blurb):
    ws.cell(2, 2, heading).font = H1
    ws.cell(3, 2, blurb).font = NOTE
    ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=4)


# ===========================================================================
# Read me
# ===========================================================================
ws = sheet("Read me")
ws.column_dimensions["B"].width = 100
ws.cell(2, 2, "Sample-Size and Power Calculator").font = Font(
    name="Calibri", size=18, bold=True, color=NAVY)
ws.cell(3, 2, "Modern Biostatistics for the Health Sciences, Volume I  "
              "—  Chapter 11, Section 11.6").font = Font(
    name="Calibri", size=11, color=GREY)

LINES = [
    "",
    ("WHAT THIS IS", H2),
    ("Five sample-size and power calculations for t-test designs, worked "
     "entirely in worksheet", LBL),
    ("formulas. The same five appear in R, SAS, SPSS, and Python in this "
     "folder, and all five", LBL),
    ("languages give the same answers.", LBL),
    "",
    ("HOW TO USE IT", H2),
    ("Each numbered sheet is one calculation. Change the values in the blue "
     "boxes; everything", LBL),
    ("below them recomputes. There are no macros, so nothing needs to be "
     "enabled, and the file", LBL),
    ("works in Excel, LibreOffice Calc, Google Sheets, and Numbers.", LBL),
    "",
    ("THE ONE THING EXCEL CANNOT DO", H2),
    ("Power calculations need the NONCENTRAL t distribution. Excel has no "
     "function for it.", LBL),
    ("R, SAS, SPSS, and Python all do. So these sheets use a standard "
     "approximation: they", LBL),
    ("replace the noncentral t with an ordinary t shifted sideways, and "
     "they chase the answer", LBL),
    ("through six refinement rows because the degrees of freedom depend on "
     "the sample size and", LBL),
    ("the sample size depends on the degrees of freedom.", LBL),
    "",
    ("The approximation is off by about one hundredth of a participant and "
     "one ten-thousandth", LBL),
    ("of a unit of power. In all five designs here it rounds to exactly the "
     "same number of", LBL),
    ("people as the exact calculation. The 'Summary' sheet shows both side "
     "by side.", LBL),
    "",
    ("That is worth knowing in itself: the third decimal place of a power "
     "calculation is not", LBL),
    ("where the risk in a study design lives. The assumed effect size is.", LBL),
    "",
    ("WHEN NOT TO USE THIS SHEET", H2),
    ("For a number that goes into a grant application, a protocol, or an "
     "IRB submission, use", LBL),
    ("R, SAS, SPSS, or Python and record the exact inputs. Use this "
     "workbook to understand", LBL),
    ("what those programs are doing, to sanity-check a number somebody "
     "hands you, and to see", LBL),
    ("instantly how the answer moves when an assumption changes.", LBL),
    "",
    ("FUNCTIONS USED", H2),
    ("T.INV.2T(alpha, df)     the two-tailed critical value of t", MONO),
    ("T.INV(p, df)            the t value with left-tail area p", MONO),
    ("T.DIST(x, df, TRUE)     the left-tail area below x", MONO),
    ("NORM.S.INV(p)           the standard normal quantile (starting value)", MONO),
    "",
    ("In LibreOffice and Google Sheets these have the same names. In Excel "
     "2007 and earlier", LBL),
    ("they were TINV, TDIST, and NORMSINV, which take their arguments "
     "differently.", LBL),
]
r = 5
for item in LINES:
    if item == "":
        r += 1
        continue
    text, font = item
    c = ws.cell(r, 2, text)
    c.font = font
    r += 1

# ===========================================================================
# helper: the refinement ladder
# ===========================================================================
def ladder(ws, row0, k, delta_ref, sd_ref, alpha_ref, power_ref):
    """Write the six-row refinement ladder. Returns (last_n_cell, next_row)."""
    ws.cell(row0, 2, "Refinement").font = H2
    ws.cell(row0, 3, "n").font = H2
    ws.cell(row0, 4, "degrees of freedom used").font = H2
    for c in (2, 3, 4):
        ws.cell(row0, c).fill = FILL_HDR
        ws.cell(row0, c).border = BOX
    kf = f"{k}*" if k == 2 else ""
    # starting value from the normal approximation
    r = row0 + 1
    ws.cell(r, 2, "start (normal)").font = LBL
    c = ws.cell(r, 3, f"={kf}(NORM.S.INV(1-{alpha_ref}/2)+NORM.S.INV("
                      f"{power_ref}))^2*{sd_ref}^2/{delta_ref}^2")
    c.font = OUT
    c.number_format = "0.0000"
    ws.cell(r, 4, "no degrees of freedom yet — uses z").font = NOTE
    prev = f"C{r}"
    for i in range(1, 7):
        r += 1
        ws.cell(r, 2, f"pass {i}").font = LBL
        df = f"MAX(ROUND(({prev}-1)*{k},0),1)"
        c = ws.cell(r, 3, f"={kf}(T.INV.2T({alpha_ref},{df})+T.INV("
                          f"{power_ref},{df}))^2*{sd_ref}^2/{delta_ref}^2")
        c.font = OUT
        c.number_format = "0.0000"
        d = ws.cell(r, 4, f"={df}")
        d.font = MONO
        prev = f"C{r}"
    ws.cell(row0 + 7, 4, "the value stops moving — that is the answer").font = NOTE
    return prev, r + 2


# ===========================================================================
# 1. One-sample n
# ===========================================================================
ws = sheet("1 One-sample n")
title_block(ws, 1, "1.  One-sample t test — how many participants?",
            "Detect a 5-unit shift from the null value when the standard "
            "deviation is 25.")
r = 5
ws.cell(r, 2, "Inputs").font = H2
ws.cell(r, 2).fill = FILL_HDR
ws.cell(r, 2).border = BOX
ws.cell(r, 3).fill = FILL_HDR
ws.cell(r, 3).border = BOX
r += 1
r = put(ws, r, "Effect to detect  (delta)", 5, "the smallest difference "
        "worth finding — a clinical judgement, not a statistical one",
        vfont=INP, fill=FILL_IN)
r = put(ws, r, "Standard deviation  (sd)", 25, "from pilot data or the "
        "literature", vfont=INP, fill=FILL_IN)
r = put(ws, r, "Significance level  (alpha)", 0.05, "two-sided",
        vfont=INP, fill=FILL_IN, fmt="0.00")
r = put(ws, r, "Target power", 0.80, "0.80 is conventional; 0.90 for a "
        "confirmatory trial", vfont=INP, fill=FILL_IN, fmt="0.00")
D, S, A, P = "$C$6", "$C$7", "$C$8", "$C$9"
r += 1
last, r = ladder(ws, r, 1, D, S, A, P)
r = put(ws, r, "Sample size (exact root)", f"={last}", None,
        font=H2, vfont=ANS, fill=FILL_ANS, fmt="0.0000")
r = put(ws, r, "PARTICIPANTS TO ENROL", f"=ROUNDUP({last},0)",
        "always round up — rounding down loses power",
        font=H2, vfont=ANS, fill=FILL_ANS, fmt="0")
r += 1
ws.cell(r, 2, "R gives 198.1513, so 199. Same answer.").font = NOTE

# ===========================================================================
# 2. Power at n
# ===========================================================================
ws = sheet("2 Power at n")
title_block(ws, 2, "2.  One-sample t test — what power does the budget buy?",
            "Same design, but recruitment is capped at 180. Solve for power "
            "instead of sample size.")
r = 5
r = put(ws, r, "Inputs", None, None, font=H2)
r = put(ws, r, "Sample size available  (n)", 180, "what the budget or the "
        "clinic will actually yield", vfont=INP, fill=FILL_IN)
r = put(ws, r, "Effect to detect  (delta)", 5, None, vfont=INP, fill=FILL_IN)
r = put(ws, r, "Standard deviation  (sd)", 25, None, vfont=INP, fill=FILL_IN)
r = put(ws, r, "Significance level  (alpha)", 0.05, None, vfont=INP,
        fill=FILL_IN, fmt="0.00")
r += 1
r = put(ws, r, "Working", None, None, font=H2)
r = put(ws, r, "Standard error  sd/sqrt(n)", "=$C$8/SQRT($C$6)", None,
        fmt="0.00000")
se = f"C{r-1}"
r = put(ws, r, "Noncentrality  delta/SE", f"=$C$7/{se}",
        "how many standard errors the true effect sits from the null",
        fmt="0.00000")
ncp = f"C{r-1}"
r = put(ws, r, "Degrees of freedom  n-1", "=$C$6-1", None, fmt="0")
df = f"C{r-1}"
r = put(ws, r, "Critical value  t(alpha/2, df)", f"=T.INV.2T($C$9,{df})",
        "the threshold the test actually uses", fmt="0.00000")
crit = f"C{r-1}"
r += 1
r = put(ws, r, "POWER", f"=1-T.DIST({crit}-{ncp},{df},TRUE)",
        "the chance this study finds the effect if it is really there",
        font=H2, vfont=ANS, fill=FILL_ANS, fmt="0.0000")
pw = f"C{r-1}"
r = put(ws, r, "Type II error  (beta)", f"=1-{pw}",
        "the chance it misses", fmt="0.0000")
r += 1
ws.cell(r, 2, "R gives 0.7608. This sheet gives 0.7607 — the "
              "approximation costs one ten-thousandth.").font = NOTE
r += 1
ws.cell(r, 2, "Below 0.80. Either enrol more people, accept a smaller "
              "power, or look for a larger effect.").font = NOTE

# ===========================================================================
# 3. Two-sample unequal
# ===========================================================================
ws = sheet("3 Two-sample unequal")
title_block(ws, 3, "3.  Two independent groups — unequal variances",
            "The groups are not expected to vary equally: sd 15 and 18. "
            "Detect a difference of 10.")
r = 5
r = put(ws, r, "Inputs", None, None, font=H2)
r = put(ws, r, "Standard deviation, group 1", 15, None, vfont=INP, fill=FILL_IN)
r = put(ws, r, "Standard deviation, group 2", 18, None, vfont=INP, fill=FILL_IN)
r = put(ws, r, "Difference to detect  (delta)", 10, None, vfont=INP, fill=FILL_IN)
r = put(ws, r, "Significance level  (alpha)", 0.05, None, vfont=INP,
        fill=FILL_IN, fmt="0.00")
r = put(ws, r, "Target power", 0.80, None, vfont=INP, fill=FILL_IN, fmt="0.00")
r += 1
r = put(ws, r, "Combined standard deviation",
        "=SQRT(($C$6^2+$C$7^2)/2)",
        "the ROOT-MEAN-SQUARE of the two, not their average (16.5)",
        font=H2, fmt="0.00000")
SD = f"$C${r-1}"
D, A, P = "$C$8", "$C$9", "$C$10"
r += 1
last, r = ladder(ws, r, 2, D, SD, A, P)
r = put(ws, r, "n per group (exact root)", f"={last}", None,
        font=H2, vfont=ANS, fill=FILL_ANS, fmt="0.0000")
r = put(ws, r, "ENROL PER GROUP", f"=ROUNDUP({last},0)", None,
        font=H2, vfont=ANS, fill=FILL_ANS, fmt="0")
g = f"C{r-1}"
r = put(ws, r, "TOTAL", f"=2*{g}", None, font=H2, vfont=ANS,
        fill=FILL_ANS, fmt="0")
r += 1
ws.cell(r, 2, "R gives 44.0716, so 45 per group. Same answer.").font = NOTE

# ===========================================================================
# 4. Two-sample equal
# ===========================================================================
ws = sheet("4 Two-sample equal")
title_block(ws, 4, "4.  Two independent groups — equal variances",
            "The same comparison assuming one common standard deviation of "
            "17. Compare with sheet 3.")
r = 5
r = put(ws, r, "Inputs", None, None, font=H2)
r = put(ws, r, "Common standard deviation", 17, None, vfont=INP, fill=FILL_IN)
r = put(ws, r, "Difference to detect  (delta)", 10, None, vfont=INP, fill=FILL_IN)
r = put(ws, r, "Significance level  (alpha)", 0.05, None, vfont=INP,
        fill=FILL_IN, fmt="0.00")
r = put(ws, r, "Target power", 0.80, None, vfont=INP, fill=FILL_IN, fmt="0.00")
D, SD, A, P = "$C$7", "$C$6", "$C$8", "$C$9"
r += 1
last, r = ladder(ws, r, 2, D, SD, A, P)
r = put(ws, r, "n per group (exact root)", f"={last}", None,
        font=H2, vfont=ANS, fill=FILL_ANS, fmt="0.0000")
r = put(ws, r, "ENROL PER GROUP", f"=ROUNDUP({last},0)", None,
        font=H2, vfont=ANS, fill=FILL_ANS, fmt="0")
g = f"C{r-1}"
r = put(ws, r, "TOTAL", f"=2*{g}", None, font=H2, vfont=ANS,
        fill=FILL_ANS, fmt="0")
r += 1
ws.cell(r, 2, "R gives 46.3467, so 47 per group. Same answer.").font = NOTE
r += 1
ws.cell(r, 2, "Two more people per group than sheet 3, for a combined "
              "spread only half a unit larger.").font = NOTE

# ===========================================================================
# 5. Paired
# ===========================================================================
ws = sheet("5 Paired")
title_block(ws, 5, "5.  Paired t test — two measurements on the same person",
            "Pairing helps exactly as much as the two measurements are "
            "correlated. That correlation has to be assumed.")
r = 5
r = put(ws, r, "Inputs", None, None, font=H2)
r = put(ws, r, "Standard deviation, measure 1", 40, None, vfont=INP, fill=FILL_IN)
r = put(ws, r, "Standard deviation, measure 2", 32, None, vfont=INP, fill=FILL_IN)
r = put(ws, r, "Correlation between them  (rho)", 0.5,
        "the assumption that matters most — see sheet 6",
        vfont=INP, fill=FILL_IN, fmt="0.00")
r = put(ws, r, "Difference to detect  (delta)", 20, None, vfont=INP, fill=FILL_IN)
r = put(ws, r, "Significance level  (alpha)", 0.05, None, vfont=INP,
        fill=FILL_IN, fmt="0.00")
r = put(ws, r, "Target power", 0.80, None, vfont=INP, fill=FILL_IN, fmt="0.00")
r += 1
r = put(ws, r, "Standard deviation of the differences",
        "=SQRT($C$6^2+$C$7^2-2*$C$8*$C$6*$C$7)",
        "shrinks as rho rises — this is the whole advantage of pairing",
        font=H2, fmt="0.00000")
SD = f"$C${r-1}"
D, A, P = "$C$9", "$C$10", "$C$11"
r += 1
last, r = ladder(ws, r, 1, D, SD, A, P)
r = put(ws, r, "Pairs (exact root)", f"={last}", None,
        font=H2, vfont=ANS, fill=FILL_ANS, fmt="0.0000")
r = put(ws, r, "PAIRS TO ENROL", f"=ROUNDUP({last},0)", None,
        font=H2, vfont=ANS, fill=FILL_ANS, fmt="0")
r += 1
ws.cell(r, 2, "R gives 28.3529, so 29 pairs. Same answer.").font = NOTE

# ===========================================================================
# 6. Rho sweep
# ===========================================================================
ws = sheet("6 Rho sweep")
ws.column_dimensions["B"].width = 14
for col, w in zip("CDEFGH", (18, 14, 14, 14, 14, 40)):
    ws.column_dimensions[col].width = w
title_block(ws, 6, "6.  How much does the assumed correlation matter?",
            "The same paired design at five values of rho. This is the "
            "sensitivity analysis a reviewer will ask for.")
r = 5
r = put(ws, r, "Fixed inputs", None, None, font=H2)
r = put(ws, r, "sd, measure 1", 40, None, vfont=INP, fill=FILL_IN)
r = put(ws, r, "sd, measure 2", 32, None, vfont=INP, fill=FILL_IN)
r = put(ws, r, "delta", 20, None, vfont=INP, fill=FILL_IN)
r = put(ws, r, "alpha", 0.05, None, vfont=INP, fill=FILL_IN, fmt="0.00")
r = put(ws, r, "power", 0.80, None, vfont=INP, fill=FILL_IN, fmt="0.00")
r += 1

hdr = ["rho", "sd of differences", "start (z)", "pass 1", "pass 2",
       "pass 3", "PAIRS"]
for j, h in enumerate(hdr):
    c = ws.cell(r, 2 + j, h)
    c.font = H2
    c.fill = FILL_HDR
    c.border = BOX
    c.alignment = Alignment(horizontal="center")
top = r + 1
for i, rho in enumerate((0.0, 0.3, 0.5, 0.7, 0.9)):
    rr = top + i
    c = ws.cell(rr, 2, rho)
    c.font = INP
    c.fill = FILL_IN
    c.border = BOX
    c.number_format = "0.0"
    c.alignment = Alignment(horizontal="right")
    ws.cell(rr, 3, f"=SQRT($C$6^2+$C$7^2-2*B{rr}*$C$6*$C$7)").number_format = "0.0000"
    ws.cell(rr, 4, f"=(NORM.S.INV(1-$C$9/2)+NORM.S.INV($C$10))^2*"
                   f"C{rr}^2/$C$8^2").number_format = "0.0000"
    prev = f"D{rr}"
    for j, col in enumerate("EFG"):
        dfx = f"MAX(ROUND({prev}-1,0),1)"
        ws.cell(rr, 5 + j, f"=(T.INV.2T($C$9,{dfx})+T.INV($C$10,{dfx}))^2*"
                           f"C{rr}^2/$C$8^2").number_format = "0.0000"
        prev = f"{col}{rr}"
    c = ws.cell(rr, 8, f"=ROUNDUP({prev},0)")
    c.font = ANS
    c.fill = FILL_ANS
    c.border = BOX
    c.number_format = "0"
    c.alignment = Alignment(horizontal="center")
    for col in range(3, 8):
        ws.cell(rr, col).font = OUT
        ws.cell(rr, col).alignment = Alignment(horizontal="right")

r = top + 6
ws.cell(r, 2, "R gives 54, 39, 29, 19, 9. Same answers.").font = NOTE
r += 2
for line in [
    "Read the last column downwards. Assuming rho = 0.7 when the truth is "
    "0.3 leaves the study with 19 pairs",
    "where it needed 39 — less than half the sample it required, and no "
    "way to tell from the data that",
    "anything went wrong. When rho is uncertain, plan on the low end. A "
    "study that is larger than it",
    "needed to be is an inefficiency; one that is smaller than it needed to "
    "be answers nothing.",
]:
    ws.cell(r, 2, line).font = NOTE
    r += 1

# ===========================================================================
# Summary
# ===========================================================================
ws = sheet("Summary")
ws.column_dimensions["B"].width = 38
for col, w in zip("CDEF", (16, 16, 16, 44)):
    ws.column_dimensions[col].width = w
ws.cell(2, 2, "All five answers, and how close the approximation comes").font = H1
ws.cell(3, 2, "The 'this workbook' column is live: it reads the other "
              "sheets. The 'R, SAS, SPSS, Python' column is typed in.").font = NOTE

r = 5
for j, h in enumerate(["Calculation", "This workbook", "Exact", "Enrol",
                       "Comment"]):
    c = ws.cell(r, 2 + j, h)
    c.font = H2
    c.fill = FILL_HDR
    c.border = BOX

ROWS = [
    ("1. One-sample, n", "='1 One-sample n'!C20", 198.1513,
     "='1 One-sample n'!C21", "199 either way"),
    ("2. Power at n = 180", "='2 Power at n'!C17", 0.7608,
     "", "agrees to three decimals"),
    ("3. Two-sample, sd 15 and 18", "='3 Two-sample unequal'!C23", 44.0716,
     "='3 Two-sample unequal'!C24", "45 per group either way"),
    ("4. Two-sample, sd 17", "='4 Two-sample equal'!C20", 46.3467,
     "='4 Two-sample equal'!C21", "47 per group either way"),
    ("5. Paired, rho = 0.5", "='5 Paired'!C24", 28.3529,
     "='5 Paired'!C25", "29 pairs either way"),
]
for i, (lbl, live, exact, enrol, note) in enumerate(ROWS):
    rr = 6 + i
    ws.cell(rr, 2, lbl).font = LBL
    c = ws.cell(rr, 3, live)
    c.font = OUT
    c.number_format = "0.0000"
    c.alignment = Alignment(horizontal="right")
    c = ws.cell(rr, 4, exact)
    c.font = OUT
    c.number_format = "0.0000"
    c.alignment = Alignment(horizontal="right")
    if enrol:
        c = ws.cell(rr, 5, enrol)
        c.font = ANS
        c.fill = FILL_ANS
        c.border = BOX
        c.number_format = "0"
        c.alignment = Alignment(horizontal="center")
    ws.cell(rr, 6, note).font = NOTE

r = 13
for line in [
    "The two middle columns never differ by enough to change the number of "
    "people you would enrol.",
    "",
    "That is the useful lesson. A power calculation is not a precision "
    "instrument and was never",
    "meant to be one. Its output is dominated by the effect size you "
    "assumed, which is a clinical",
    "judgement, and by the standard deviation you borrowed from somebody "
    "else's study. Arguing about",
    "the fourth decimal place of the arithmetic while those two numbers are "
    "guesses is misplaced effort.",
    "",
    "Report the assumptions. Report where they came from. Report what "
    "happens if they are wrong.",
    "That is what sheet 6 is for, and it is what a reviewer actually reads.",
]:
    ws.cell(r, 2, line).font = NOTE if line else LBL
    r += 1

del wb["Sheet"]
wb.active = 0

# ---------------------------------------------------------------------------
# Excel 2010 introduced T.INV, T.INV.2T, T.DIST and NORM.S.INV. In the OOXML
# file format every post-2007 function must be stored with an "_xlfn." prefix;
# Excel strips it on display. openpyxl does not add it, so without this pass
# every formula opens as #NAME?.
# ---------------------------------------------------------------------------
MODERN = ("NORM.S.INV", "T.INV.2T", "T.INV", "T.DIST")


def prefix(formula):
    # Longest names first, so T.INV.2T( is consumed before T.INV( is tried.
    # A marker byte stands in for the prefix so no name is rewritten twice.
    out = formula
    for fn in MODERN:
        out = out.replace(fn + "(", "\x00" + fn + "(")
    return out.replace("\x00", "_xlfn.")


for ws in wb.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.value = prefix(cell.value)

wb.save("/home/claude/ssz/code/lab_sample_size.xlsx")
print("written: lab_sample_size.xlsx")
print("sheets:", wb.sheetnames)
